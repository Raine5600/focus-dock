#!/usr/bin/env python3
"""Build affiliate outreach workbooks organized by contact_type."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INPUT = Path(__file__).parent.parent / "marketing" / "affiliate-contacts-master.json"
OUT_XLSX = Path(__file__).parent.parent / "marketing" / "affiliate-contacts-by-type.xlsx"
OUT_JSON = Path(__file__).parent.parent / "marketing" / "affiliate-contacts-master.json"

HEADERS = [
    "channel_name",
    "contact_name",
    "platform",
    "profile_url",
    "subscribers_or_audience",
    "niche_focus",
    "contact_email",
    "contact_type",
    "contact_url",
    "sells_digital_products",
    "outreach_priority",
    "notes",
]

HEADER_FILL = PatternFill("solid", fgColor="1a1f3d")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial")
WIDTHS = [28, 16, 12, 42, 14, 32, 28, 14, 42, 8, 10, 40]


def load_contacts():
    with open(INPUT) as f:
        data = json.load(f)
    valid = []
    for c in data:
        ct = (c.get("contact_type") or "").strip()
        url = (c.get("contact_url") or "").strip()
        email = (c.get("contact_email") or "").strip()
        if ct not in ("email", "contact_form"):
            continue
        if ct == "email" and not email:
            continue
        if not url and ct == "contact_form":
            continue
        if ct == "email" and not url:
            c["contact_url"] = f"mailto:{email}"
        valid.append(c)
    return valid


def write_sheet(ws, rows):
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row_idx, c in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=c.get("channel_name", ""))
        ws.cell(row=row_idx, column=2, value=c.get("contact_name", ""))
        ws.cell(row=row_idx, column=3, value=c.get("platform", ""))
        ws.cell(row=row_idx, column=4, value=c.get("profile_url", ""))
        ws.cell(row=row_idx, column=5, value=c.get("subscribers_or_audience", c.get("subscribers", "")))
        ws.cell(row=row_idx, column=6, value=c.get("niche_focus", ""))
        ws.cell(row=row_idx, column=7, value=c.get("contact_email", ""))
        ws.cell(row=row_idx, column=8, value=c.get("contact_type", ""))
        ws.cell(row=row_idx, column=9, value=c.get("contact_url", ""))
        ws.cell(row=row_idx, column=10, value=c.get("sells_digital_products", ""))
        ws.cell(row=row_idx, column=11, value=c.get("outreach_priority", ""))
        ws.cell(row=row_idx, column=12, value=c.get("notes", ""))
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def main():
    contacts = load_contacts()
    contacts.sort(key=lambda x: (x.get("contact_type", ""), x.get("outreach_priority", "z"), x.get("channel_name", "")))

    with open(OUT_JSON, "w") as f:
        json.dump(contacts, f, indent=2)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    by_type = {}
    for c in contacts:
        by_type.setdefault(c["contact_type"], []).append(c)
    summary.append(["contact_type", "count"])
    for t in sorted(by_type):
        summary.append([t, len(by_type[t])])
    summary.append(["TOTAL", len(contacts)])

    all_sheet = wb.create_sheet("All Contacts")
    write_sheet(all_sheet, contacts)

    for t in sorted(by_type):
        name = f"email" if t == "email" else "contact_form"
        ws = wb.create_sheet(name[:31])
        write_sheet(ws, by_type[t])

    wb.save(OUT_XLSX)
    print(f"Saved {len(contacts)} contacts")
    for t, rows in sorted(by_type.items()):
        print(f"  {t}: {len(rows)}")
    print(f"JSON: {OUT_JSON}")
    print(f"XLSX: {OUT_XLSX}")


if __name__ == "__main__":
    main()