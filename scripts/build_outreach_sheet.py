#!/usr/bin/env python3
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MARKETING_DIR = Path(__file__).resolve().parent.parent / "marketing"
JSON_PATH = MARKETING_DIR / "creator-outreach-list.json"
XLSX_PATH = MARKETING_DIR / "creator-outreach-list.xlsx"

with open(JSON_PATH) as f:
    creators = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "Creator Outreach"

headers = [
    "channel_name",
    "contact_name",
    "youtube_url",
    "subscribers",
    "niche_focus",
    "contact_email",
    "contact_type",
    "contact_url",
    "sells_digital_products",
    "outreach_priority",
    "notes",
    "personalized_outreach_note",
    "outreach_status",
    "email_sent_date",
]
header_fill = PatternFill("solid", fgColor="1a1f3d")
header_font = Font(bold=True, color="FFFFFF", name="Arial")

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for row, c in enumerate(creators, 2):
    ws.cell(row=row, column=1, value=c["channel_name"])
    ws.cell(row=row, column=2, value=c.get("contact_name", ""))
    ws.cell(row=row, column=3, value=c["youtube_url"])
    ws.cell(row=row, column=4, value=c["subscribers"])
    ws.cell(row=row, column=5, value=c["niche_focus"])
    ws.cell(row=row, column=6, value=c["contact_email"])
    ws.cell(row=row, column=7, value=c["contact_type"])
    ws.cell(row=row, column=8, value=c["contact_url"])
    ws.cell(row=row, column=9, value=c["sells_digital_products"])
    ws.cell(row=row, column=10, value=c["outreach_priority"])
    ws.cell(row=row, column=11, value=c["notes"])
    ws.cell(row=row, column=12, value=c.get("personalized_outreach_note", ""))
    ws.cell(row=row, column=13, value="pending")
    ws.cell(row=row, column=14, value="")

widths = [28, 14, 42, 12, 35, 30, 14, 42, 10, 10, 50, 55, 14, 14]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(XLSX_PATH)
print(f"Saved {XLSX_PATH} with {len(creators)} creators")